"""Tests use isolated copies of the supplied template, never the live workbook."""
from copy import deepcopy
from datetime import date, datetime
import hashlib
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook

from services.excel_service import ExcelService, ORDER_HEADERS


@pytest.fixture
def report(tmp_path):
    root = Path(__file__).resolve().parents[1]
    original = root / "backups" / "original_template.xlsx"
    source = original if original.exists() else root / "EasyOats_Order_Tracker.xlsx"
    workbook = tmp_path / "EasyOats_Order_Tracker.xlsx"
    shutil.copy2(source, workbook)
    return ExcelService(workbook, tmp_path / "backups")


def sample_snapshot(count=1):
    orders = []
    for index in range(count):
        orders.append({
            "order_id": f"EO-{index + 1:06}", "customer_name": "عميل تجريبي للاختبار",
            "phone_original": "+20 1012345678", "phone_normalized": "01012345678",
            "order_datetime": datetime(2026, 9, 16, 10, 30), "source": "واتساب",
            "area": "القاهرة", "address": "العنوان", "location_url": "https://maps.google.com/",
            "order_type": "سعر عادي", "honey_qty": 1, "date_qty": 1,
            "discount": 5, "delivery_fee": 25, "delivery_cost": 15,
            "payment_method": "إنستاباي", "amount_collected": 200,
            "status": "جديد", "courier": "فريق EasyOats", "notes": "=HYPERLINK(\"bad\",\"bad\")",
            "feedback_consent": "موافق", "feedback_status": "لم يُطلب", "returned_sellable": False,
            "refund_amount": 0, "total_units": 2, "unit_price": 75, "product_subtotal": 150,
            "total_due": 170, "product_cost": 88, "contribution_margin": 67,
            "outstanding_balance": -30, "payment_status": "مدفوع", "next_action": "لا يوجد إجراء عاجل",
            "honey_unit_cost": 44, "date_unit_cost": 44, "customer_rating": 4, "buy_again": "نعم",
        })
    return {
        "orders": orders, "feedback": [],
        "inventory": [dict(sku=sku, name=name, opening_stock=1000, added_stock=0,
            reserved=count, delivered=0, returned_unsellable=0, available=1000-count,
            physical_count=None, variance=None, reorder_point=50, unit_cost=44,
            value=(1000-count)*44, low_stock=False)
            for sku, name in (("honey", "عسل ولبن"), ("date", "دبس تمر ولبن"))],
        "settings": dict(retail_price=75, offer_price=120, honey_unit_cost=44, date_unit_cost=44,
            low_stock_threshold=50, google_form_url="https://forms.google.com/test", current_user="فريق EasyOats", user_names=["فريق EasyOats"]),
    }


def column(ws, label):
    return next(cell.column for cell in ws[6] if cell.value == label)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_backup_is_byte_identical_and_refresh_preserves_original(report):
    before = digest(report.workbook_path)
    first = report.sync(sample_snapshot())
    assert first.success, first.message
    assert first.timestamp is not None
    original = report.backup_dir / "original_template.xlsx"
    assert digest(original) == before
    assert len(list(report.backup_dir.glob("EasyOats_*.xlsx"))) == 1
    snapshot = sample_snapshot()
    snapshot["orders"][0]["notes"] = "تحديث"
    assert report.sync(snapshot).success
    assert digest(original) == before
    assert len(list(report.backup_dir.glob("EasyOats_*.xlsx"))) == 2


def test_missing_operational_workbook_is_initialized_from_read_only_template(tmp_path):
    root = Path(__file__).resolve().parents[1]
    target = tmp_path / "persistent" / "EasyOats_Order_Tracker.xlsx"
    service = ExcelService(target, tmp_path / "persistent" / "backups")
    service.template_path = root / "templates" / "EasyOats_Order_Tracker.xlsx"
    expected = digest(service.template_path)
    assert service.sync(sample_snapshot()).success
    assert target.exists()
    assert digest(service.backup_dir / "original_template.xlsx") == expected


def test_atomic_replace_failure_leaves_workbook_intact(report, monkeypatch):
    before = digest(report.workbook_path)
    def locked(*args, **kwargs):
        raise PermissionError("Workbook open in Excel")
    monkeypatch.setattr("services.excel_service.os.replace", locked)
    result = report.sync(sample_snapshot())
    assert not result.success
    assert "أغلق" in result.message
    assert digest(report.workbook_path) == before
    assert not list(report.workbook_path.parent.glob(".EasyOats_*.tmp.xlsx"))
    assert digest(report.backup_dir / "original_template.xlsx") == before


def test_more_than_200_preserves_formulas_charts_formats_dropdowns(report):
    snapshot = sample_snapshot(240)
    assert report.sync(snapshot).success
    workbook = load_workbook(report.workbook_path)
    orders = workbook["الطلبات"]
    assert orders["A246"].value == "EO-000240"
    assert orders["K246"].data_type == "f"
    assert orders["K246"].style_id == orders["K8"].style_id
    assert orders.tables["OrdersTable"].ref.endswith("246")
    assert all("246" in str(item.sqref) for item in orders.data_validations.dataValidation)
    assert all("246" in str(item.sqref) for item in orders.conditional_formatting)
    assert len(workbook["لوحة المتابعة"]._charts) == 1
    assert "$246" in workbook["لوحة المتابعة"]["A5"].value
    assert "$246" in workbook["المخزون"]["E6"].value
    assert "$246" in workbook["الفيدباك"]["D7"].value
    assert workbook.sheetnames == ["لوحة المتابعة", "الطلبات", "المخزون", "الفيدباك", "القوائم والإعدادات"]
    workbook.close()
    cached = load_workbook(report.workbook_path, data_only=True)
    assert cached["لوحة المتابعة"]["A5"].value == 240
    assert cached["الطلبات"]["W246"].value == -30
    assert cached["المخزون"]["G6"].value == 760
    cached.close()


def test_columns_follow_headers_after_reordering(report):
    workbook = load_workbook(report.workbook_path)
    orders = workbook["الطلبات"]
    # Move customer name and phone, and move the quantity used by formulas.
    for left, right in ((4, 5), (9, 10)):
        for row in range(6, orders.max_row + 1):
            orders.cell(row, left).value, orders.cell(row, right).value = orders.cell(row, right).value, orders.cell(row, left).value
    workbook.save(report.workbook_path)
    workbook.close()
    snapshot = sample_snapshot()
    snapshot["orders"][0]["honey_qty"] = 3
    assert report.sync(snapshot).success
    workbook = load_workbook(report.workbook_path)
    orders = workbook["الطلبات"]
    assert orders["E7"].value == "عميل تجريبي للاختبار"
    assert orders["D7"].value == "+20 1012345678"
    assert orders["D7"].data_type == "s"
    assert orders["J7"].value == 3
    assert "J7+I7" in orders["K7"].value
    assert "$J$7" in workbook["المخزون"]["E6"].value
    assert "$E$7" in workbook["الفيدباك"]["D7"].value
    workbook.close()


def test_frozen_prices_unclamped_balance_and_formula_injection(report):
    snapshot = sample_snapshot()
    snapshot["settings"]["retail_price"] = 99
    snapshot["settings"]["honey_unit_cost"] = 60
    assert report.sync(snapshot).success
    workbook = load_workbook(report.workbook_path)
    orders = workbook["الطلبات"]
    assert "MAX(" not in orders["Q7"].value
    assert "MAX(" not in orders["W7"].value
    assert "القوائم" not in orders["L7"].value
    assert "المخزون" not in orders["R7"].value
    assert orders["AI7"].data_type == "s"
    assert orders["AI7"].value.startswith("=HYPERLINK")
    assert orders.cell(7, column(orders, ORDER_HEADERS["phone_normalized"])).value == "01012345678"
    workbook.close()
    workbook = load_workbook(report.workbook_path, data_only=True)
    orders = workbook["الطلبات"]
    assert orders["L7"].value == 75
    assert orders["Q7"].value == 170
    assert orders["R7"].value == 88
    assert orders["W7"].value == -30
    assert orders["U7"].value == "مدفوع"
    assert orders["AG7"].value == 4
    workbook.close()


def test_returned_unsellable_inventory_formula(report):
    snapshot = sample_snapshot()
    snapshot["orders"][0]["status"] = "مرتجع"
    for item in snapshot["inventory"]:
        item.update(reserved=0, returned_unsellable=1)
    assert report.sync(snapshot).success
    workbook = load_workbook(report.workbook_path)
    assert "-N6" in workbook["المخزون"]["G6"].value
    assert '"مرتجع"' in workbook["المخزون"]["N6"].value
    assert '"<>نعم"' in workbook["المخزون"]["N6"].value
    workbook.close()


def test_refuses_to_erase_existing_orders_if_database_is_empty(report):
    assert report.sync(sample_snapshot()).success
    before = digest(report.workbook_path)
    result = report.sync(sample_snapshot(0))
    assert not result.success
    assert "غير موجودة" in result.message
    assert digest(report.workbook_path) == before


def test_feedback_expansion_and_cached_customer_name(report):
    snapshot = sample_snapshot()
    snapshot["feedback"] = [{"id": index + 1, "order_id": "EO-000001",
        "response_date": date(2026, 9, 16), "source": "واتساب", "consent_confirmed": True,
        "overall_rating": 4, "taste_rating": 5, "portion_rating": 4,
        "preparation_rating": 5, "value_rating": 4, "buy_again": "نعم",
        "recommendation_score": 9, "preferred_flavor": "عسل ولبن",
        "comments": "+external", "issue_reported": False, "followup_status": "لا تحتاج"}
        for index in range(205)]
    assert report.sync(snapshot).success
    workbook = load_workbook(report.workbook_path)
    sheet = workbook["الفيدباك"]
    assert sheet.tables["FeedbackTable"].ref == "A6:S211"
    assert sheet["D211"].data_type == "f"
    assert sheet["O211"].data_type == "s"
    # Order summaries are authoritative, including manual corrections made
    # after feedback was received; the formula therefore reads the stored field.
    assert "AR7" in workbook["الطلبات"]["AG7"].value
    assert "'الطلبات'!$AG$7:$AG$206" in workbook["لوحة المتابعة"]["E11"].value
    workbook.close()
    workbook = load_workbook(report.workbook_path, data_only=True)
    assert workbook["الفيدباك"]["D211"].value == "عميل تجريبي للاختبار"
    assert workbook["الفيدباك"]["F211"].value == "نعم"
    workbook.close()


def test_snapshot_callback_is_inside_cross_process_lock(report, monkeypatch):
    from filelock import FileLock
    called = []
    def snapshot():
        lock = FileLock(report.lock_path, timeout=0)
        from filelock import Timeout
        with pytest.raises(Timeout):
            lock.acquire()
        called.append(True)
        return sample_snapshot()
    assert report.sync(snapshot).success
    assert called == [True]
