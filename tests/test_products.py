from io import BytesIO

import pytest
from openpyxl import load_workbook

from services.excel_service import ORDER_HEADERS
from services.order_service import ValidationError


def product_row(service, sku):
    return next(item for item in service.inventory() if item["sku"] == sku)


def test_admin_can_add_product_and_use_it_in_orders(service, order_data):
    product = service.create_product({
        "sku": "chocolate",
        "name": "شوكولاتة ولبن",
        "opening_stock": 10,
        "unit_cost": 30,
        "reorder_point": 3,
    }, user="مسؤول الاختبار", is_admin=True)
    assert product["sku"] == "chocolate"
    assert product["active"] is True

    order_data["product_quantities"] = {"honey": 1, "chocolate": 2}
    order = service.create_order(order_data)
    assert order["product_quantities"] == {"chocolate": 2, "honey": 1}
    assert order["total_units"] == 3
    assert order["product_cost"] == 104
    assert order["product_subtotal"] == 225
    assert "شوكولاتة ولبن: 2" in order["products_summary"]
    stock = product_row(service, "chocolate")
    assert stock["reserved"] == 2
    assert stock["available"] == 8


def test_product_cost_is_frozen_per_order(service, order_data):
    service.create_product({
        "sku": "banana", "name": "موز ولبن", "opening_stock": 20,
        "unit_cost": 25, "reorder_point": 4,
    }, user="مسؤول", is_admin=True)
    order_data["product_quantities"] = {"banana": 2}
    first = service.create_order(order_data)

    service.update_product_definition(
        "banana", {"unit_cost": 40}, user="مسؤول", is_admin=True,
    )
    service.update_order(first["order_id"], {"notes": "لا يغير التكلفة القديمة"})
    old = service.get_order(first["order_id"])
    second = service.create_order(order_data)
    assert old["product_cost"] == 50
    assert second["product_cost"] == 80


def test_inactive_product_remains_on_history_but_cannot_be_added(service, order_data):
    service.create_product({
        "sku": "berry", "name": "توت ولبن", "opening_stock": 10,
        "unit_cost": 35, "reorder_point": 2,
    }, user="مسؤول", is_admin=True)
    order_data["product_quantities"] = {"berry": 1}
    existing = service.create_order(order_data)
    service.update_product_definition(
        "berry", {"active": False}, user="مسؤول", is_admin=True,
    )

    service.update_order(existing["order_id"], {"notes": "يبقى في السجل"})
    assert service.get_order(existing["order_id"])["product_quantities"] == {"berry": 1}
    order_data["phone_original"] = "01098765432"
    with pytest.raises(ValidationError, match="متوقف"):
        service.create_order(order_data)


def test_product_management_requires_admin(service):
    with pytest.raises(ValidationError, match="مسؤول"):
        service.create_product({
            "sku": "plain", "name": "سادة", "opening_stock": 1,
            "unit_cost": 10, "reorder_point": 1,
        }, user="موظف")


def test_dynamic_product_is_exported_and_survives_excel_order_update(service, order_data):
    service.create_product({
        "sku": "cocoa", "name": "كاكاو ولبن", "opening_stock": 12,
        "unit_cost": 32, "reorder_point": 3,
    }, user="مسؤول", is_admin=True)
    order_data["product_quantities"] = {"honey": 1, "cocoa": 2}
    order = service.create_order(order_data)
    assert service.sync_excel()

    workbook = load_workbook(service.workbook_path, data_only=False)
    orders = workbook["الطلبات"]
    header = {cell.value: cell.column for cell in orders[6]}
    row = next(index for index in range(7, orders.max_row + 1)
               if orders.cell(index, header[ORDER_HEADERS["order_id"]]).value == order["order_id"])
    summary = orders.cell(row, header[ORDER_HEADERS["products_summary"]]).value
    assert "كاكاو ولبن: 2" in summary
    assert "عسل ولبن: 1" in summary
    assert orders.cell(row, header[ORDER_HEADERS["extra_units"]]).value == 2
    assert "extra" not in str(orders.cell(row, header[ORDER_HEADERS["total_units"]]).value).casefold()
    inventory = workbook["المخزون"]
    cocoa_row = next(index for index in range(1, inventory.max_row + 1)
                     if inventory.cell(index, 1).value == "cocoa")
    assert inventory.cell(cocoa_row, 5).value == 2
    assert inventory.cell(cocoa_row, 7).value == 10
    workbook.close()

    # An order-field import must preserve product quantities that are represented
    # by the dynamic product model rather than the two legacy flavor columns.
    workbook = load_workbook(service.workbook_path)
    orders = workbook["الطلبات"]
    header = {cell.value: cell.column for cell in orders[6]}
    row = next(index for index in range(7, orders.max_row + 1)
               if orders.cell(index, header[ORDER_HEADERS["order_id"]]).value == order["order_id"])
    orders.cell(row, header[ORDER_HEADERS["notes"]]).value = "تحديث Excel"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    preview = service.preview_excel_order_import(content)
    service.import_excel_order_updates(
        content, preview["revision"], user="مسؤول", is_admin=True,
        import_reason="اختبار الحفاظ على المنتج",
    )
    updated = service.get_order(order["order_id"])
    assert updated["product_quantities"] == {"cocoa": 2, "honey": 1}
