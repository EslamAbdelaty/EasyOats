from io import BytesIO

import pytest
from openpyxl import load_workbook

from services.excel_service import ORDER_HEADERS
from services.order_service import ValidationError


def _column(sheet, field):
    label = ORDER_HEADERS[field]
    return next(cell.column for cell in sheet[6] if cell.value == label)


def _edited_export(service, edits):
    assert service.sync_excel()
    workbook = load_workbook(BytesIO(service.workbook_path.read_bytes()))
    sheet = workbook["الطلبات"]
    rows = {
        str(sheet.cell(row, _column(sheet, "order_id")).value): row
        for row in range(7, sheet.max_row + 1)
        if sheet.cell(row, _column(sheet, "order_id")).value
    }
    for order_id, fields in edits.items():
        for field, value in fields.items():
            sheet.cell(rows[order_id], _column(sheet, field)).value = value
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_excel_import_previews_and_atomically_applies_existing_order_changes(service, order_data):
    order = service.create_order(order_data)
    content = _edited_export(service, {
        order["order_id"]: {"notes": "تحديث من ملف الفريق", "amount_collected": 50},
    })

    preview = service.preview_excel_order_import(content)
    assert preview["changed_orders"] == 1
    assert preview["change_count"] == 2
    assert {item["field"] for item in preview["changes"]} == {"notes", "amount_collected"}

    result = service.import_excel_order_updates(
        content, preview["revision"], user="مسؤول الاختبار", is_admin=True,
        import_reason="تحديث ملف متابعة الفريق",
    )
    assert result == {"changed_orders": 1, "change_count": 2}
    updated = service.get_order(order["order_id"])
    assert updated["notes"] == "تحديث من ملف الفريق"
    assert updated["amount_collected"] == 50
    audit = service.audit_history(order["order_id"])
    assert {row["field"] for row in audit if row["action"] == "excel_import"} >= {"notes", "amount_collected"}


def test_excel_import_rejects_stale_export(service, order_data):
    order = service.create_order(order_data)
    content = _edited_export(service, {order["order_id"]: {"notes": "نسخة قديمة"}})
    service.update_order(order["order_id"], {"notes": "تحديث أحدث داخل التطبيق"})

    with pytest.raises(ValidationError, match="أقدم"):
        service.preview_excel_order_import(content)
    assert service.get_order(order["order_id"])["notes"] == "تحديث أحدث داخل التطبيق"


def test_excel_import_rejects_unknown_order_without_partial_writes(service, order_data):
    first = service.create_order(order_data)
    order_data = dict(order_data)
    order_data["phone_original"] = "01098765432"
    second = service.create_order(order_data)
    content = _edited_export(service, {
        first["order_id"]: {"notes": "يجب ألا يُحفظ"},
        second["order_id"]: {"order_id": "EO-999999"},
    })

    with pytest.raises(ValidationError, match="لا يمكن إنشاء طلبات جديدة"):
        service.preview_excel_order_import(content)
    assert service.get_order(first["order_id"])["notes"] == ""


def test_excel_import_requires_current_easy_oats_export(service, order_data):
    order = service.create_order(order_data)
    content = _edited_export(service, {order["order_id"]: {"notes": "تحديث"}})
    workbook = load_workbook(BytesIO(content))
    del workbook.custom_doc_props["EasyOatsDatabaseRevision"]
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(ValidationError, match="رقم مزامنة"):
        service.preview_excel_order_import(output.getvalue())
