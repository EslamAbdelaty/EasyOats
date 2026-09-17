"""Real SQLite + template workbook integration, isolated from operational data."""
from datetime import date
from pathlib import Path
import shutil

from openpyxl import load_workbook
import pytest


@pytest.fixture
def integrated_service(tmp_path):
    from services.order_service import AppService

    source = Path(__file__).resolve().parents[1]
    original = source / "backups" / "original_template.xlsx"
    if not original.exists():
        original = source / "EasyOats_Order_Tracker.xlsx"
    target = tmp_path / "EasyOats_Order_Tracker.xlsx"
    shutil.copy2(original, target)
    app = AppService(
        database_url=f"sqlite:///{(tmp_path / 'integration.db').as_posix()}",
        workbook_path=target, backup_dir=tmp_path / "backups",
    )
    yield app, target
    if hasattr(app, "engine"):
        app.engine.dispose()


def workbook_order(path, order_id):
    with path.open("rb") as handle:
        book = load_workbook(handle, data_only=True)
        sheet = book["الطلبات"]
        headers = {cell.value: cell.column for cell in sheet[6] if cell.value}
        for row in sheet.iter_rows(min_row=7):
            if row[headers["رقم الطلب"] - 1].value == order_id:
                result = {name: row[column - 1].value for name, column in headers.items()}
                book.close()
                return result
        book.close()
    raise AssertionError(f"Missing exported order {order_id}")


def test_full_lifecycle_sqlite_and_excel(integrated_service, order_data):
    service, workbook = integrated_service
    order = service.create_order(order_data)
    assert not service.sync_status()["pending"]
    assert service.search_orders("+20 (10) 1234-5678")[0]["order_id"] == order["order_id"]
    service.update_order(order["order_id"], {"status": "مؤكد"})
    service.update_order(order["order_id"], {
        "status": "تم التوصيل", "actual_delivery_date": date(2026, 9, 2),
        "amount_collected": 185, "courier": "مسؤول الاختبار",
    })
    saved = service.get_order(order["order_id"])
    assert saved["status"] == "تم التوصيل"
    assert saved["payment_status"] == "مدفوع"
    exported = workbook_order(workbook, order["order_id"])
    assert exported["حالة الطلب"] == "تم التوصيل"
    assert exported["إجمالي المطلوب"] == 185
    assert exported["المتبقي"] == 0
    assert exported["حالة الدفع"] == "مدفوع"
    assert (workbook.parent / "backups" / "original_template.xlsx").exists()


def test_locked_export_keeps_committed_order_and_retry_works(integrated_service, order_data, monkeypatch):
    from services import excel_service

    service, workbook = integrated_service
    previous = workbook.read_bytes()
    original_replace = excel_service.os.replace

    def locked(*_args):
        raise PermissionError("Workbook open in Excel")

    monkeypatch.setattr(excel_service.os, "replace", locked)
    order = service.create_order(order_data)
    assert service.get_order(order["order_id"])["total_due"] == 185
    assert service.sync_status()["pending"]
    assert workbook.read_bytes() == previous
    monkeypatch.setattr(excel_service.os, "replace", original_replace)
    assert service.sync_excel()
    assert not service.sync_status()["pending"]
    assert workbook_order(workbook, order["order_id"])["اسم العميل"] == order_data["customer_name"]


def test_partial_refund_excel_matches_database(integrated_service, order_data):
    service, workbook = integrated_service
    order_data["amount_collected"] = 185
    order = service.create_order(order_data)
    for refund, status in [(50, "مدفوع"), (185, "مسترد")]:
        saved = service.update_order(order["order_id"], {"refund_amount": refund})
        assert saved["payment_status"] == status
        assert workbook_order(workbook, order["order_id"])["حالة الدفع"] == status
