"""Read controlled order updates from an EasyOats Excel export."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any
import unicodedata
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from services.excel_service import DEMO_IDS, ORDER_HEADERS


REVISION_PROPERTY = "EasyOatsDatabaseRevision"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 2_000

# These fields are editable in the orders sheet. Prices, costs, totals, payment
# status, balances, and next actions remain calculated by the application.
ORDER_IMPORT_FIELDS = (
    "order_datetime", "source", "customer_name", "phone_original", "area", "address",
    "order_type", "honey_qty", "date_qty", "discount", "delivery_fee", "delivery_cost",
    "payment_method", "amount_collected", "status", "courier", "expected_delivery_date",
    "actual_delivery_date", "tracking_number", "delivery_result", "feedback_consent",
    "feedback_request_date", "feedback_status", "customer_rating", "buy_again", "notes",
    "location_url", "returned_sellable", "refund_amount",
)


class ExcelImportError(ValueError):
    """A workbook validation error safe to display to the user."""


@dataclass(frozen=True)
class ParsedOrderWorkbook:
    revision: int
    rows: dict[str, dict[str, Any]]


def _normalize(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).strip().split())


def _find_orders_sheet(workbook):
    required = {_normalize(ORDER_HEADERS[key]) for key in ("order_id", "customer_name", "status")}
    for sheet in workbook.worksheets:
        for row_number in range(1, min(25, sheet.max_row) + 1):
            names = {_normalize(sheet.cell(row_number, column).value): column
                     for column in range(1, sheet.max_column + 1)
                     if sheet.cell(row_number, column).value is not None}
            if required.issubset(names):
                columns = {}
                for key in ("order_id", *ORDER_IMPORT_FIELDS, "stored_customer_rating", "stored_buy_again"):
                    title = _normalize(ORDER_HEADERS[key])
                    if title not in names:
                        raise ExcelImportError(f"عمود «{ORDER_HEADERS[key]}» غير موجود في ملف Excel.")
                    columns[key] = names[title]
                return sheet, row_number, columns
    raise ExcelImportError("لم يتم العثور على ورقة الطلبات الأصلية داخل ملف Excel.")


def _workbook_revision(workbook) -> int:
    try:
        value = workbook.custom_doc_props[REVISION_PROPERTY].value
        revision = int(value)
    except (KeyError, TypeError, ValueError, AttributeError):
        raise ExcelImportError(
            "هذا الملف لا يحتوي رقم مزامنة EasyOats. نزّل نسخة Excel جديدة من التطبيق ثم عدّلها."
        ) from None
    if revision < 0:
        raise ExcelImportError("رقم مزامنة ملف Excel غير صحيح.")
    return revision


def parse_order_workbook(content: bytes) -> ParsedOrderWorkbook:
    """Parse an exported workbook without trusting its extension or formulas."""
    if not isinstance(content, bytes) or not content:
        raise ExcelImportError("اختر ملف Excel بصيغة xlsx.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ExcelImportError("حجم ملف Excel أكبر من الحد المسموح وهو 10 ميجابايت.")
    if not content.startswith(b"PK"):
        raise ExcelImportError("الملف المرفوع ليس ملف xlsx صحيحاً.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
                raise ExcelImportError("محتوى ملف Excel أكبر من الحد الآمن للاستيراد.")
            if any(item.flag_bits & 0x1 for item in members):
                raise ExcelImportError("لا يمكن استيراد ملف Excel محمي بكلمة مرور.")
    except ExcelImportError:
        raise
    except BadZipFile:
        raise ExcelImportError("الملف المرفوع ليس ملف xlsx صحيحاً.") from None

    formulas = None
    try:
        formulas = load_workbook(BytesIO(content), data_only=False, read_only=True, keep_links=False)
        revision = _workbook_revision(formulas)
        sheet, header_row, columns = _find_orders_sheet(formulas)
        rows: dict[str, dict[str, Any]] = {}
        for row_number in range(header_row + 1, sheet.max_row + 1):
            id_cell = sheet.cell(row_number, columns["order_id"])
            if id_cell.data_type == "f":
                raise ExcelImportError(f"رقم الطلب في الصف {row_number} لا يمكن أن يكون معادلة.")
            order_id = str(id_cell.value or "").strip()
            if not order_id or order_id in DEMO_IDS:
                continue
            if order_id in rows:
                raise ExcelImportError(f"رقم الطلب {order_id} مكرر داخل ملف Excel.")
            payload: dict[str, Any] = {}
            for field in ORDER_IMPORT_FIELDS:
                cell = sheet.cell(row_number, columns[field])
                if cell.data_type == "f":
                    # Rating and repurchase are displayed through formulas backed
                    # by literal stored values in the exported workbook.
                    if field not in {"customer_rating", "buy_again"}:
                        raise ExcelImportError(
                            f"الحقل «{ORDER_HEADERS[field]}» للطلب {order_id} لا يمكن أن يكون معادلة."
                        )
                    stored_field = "stored_customer_rating" if field == "customer_rating" else "stored_buy_again"
                    value = sheet.cell(row_number, columns[stored_field]).value
                else:
                    value = cell.value
                if field == "buy_again" and value is None:
                    value = ""
                payload[field] = value
            rows[order_id] = payload
        return ParsedOrderWorkbook(revision=revision, rows=rows)
    except ExcelImportError:
        raise
    except Exception:
        raise ExcelImportError(
            "تعذرت قراءة ملف Excel. استخدم نسخة xlsx نُزّلت حديثاً من EasyOats ولا تغيّر أسماء الأوراق أو الأعمدة."
        ) from None
    finally:
        if formulas is not None:
            formulas.close()
